"""Write the sheets (E8): invoices.csv, line_items.csv, ibis.xlsx and, optionally, a Google Sheet.

All three outputs come from the same `Sheet` rows, so they always carry the same data. CSVs keep
amounts exact (Decimal as text); Excel and Sheets get real numbers and dates.

Text from invoices is never allowed to run as a formula: a description like "=HYPERLINK(...)"
is written as plain text in every output.
"""

from __future__ import annotations

import csv
import os
import re
import tempfile
from collections.abc import Callable, Iterable
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from enum import Enum
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from ibis_docs.models import Invoice, Status
from ibis_docs.results import Result

INVOICES_CSV = "invoices.csv"
LINE_ITEMS_CSV = "line_items.csv"
WORKBOOK = "ibis.xlsx"

# Brand colours: ledger green headers, amber only for "needs review".
HEADER_FILL = PatternFill("solid", fgColor="23735F")
HEADER_FONT = Font(bold=True, color="FFFFFF")
REVIEW_FILL = PatternFill("solid", fgColor="F3E6CF")  # amber #B7791F, lightened for a row

FORMULA_START = ("=", "+", "-", "@", "\t", "\r")


class Kind(Enum):
    TEXT = "text"
    MONEY = "money"
    NUMBER = "number"  # quantities, rates, counts
    DATE = "date"
    DATETIME = "datetime"


@dataclass(frozen=True)
class Column:
    header: str
    kind: Kind
    value: Callable


@dataclass
class Sheet:
    name: str
    columns: list[Column]
    rows: list[list]  # values as Python types: str, Decimal, int, date, datetime or None
    review_rows: set[int]  # row indexes to highlight

    @property
    def headers(self) -> list[str]:
        return [c.header for c in self.columns]


def _status(r: Result) -> str:
    return r.invoice.status.value


def _reasons(r: Result) -> str:
    return "; ".join(reason.message for reason in r.invoice.reasons)


def _codes(r: Result) -> str:
    return ", ".join(reason.code for reason in r.invoice.reasons)


INVOICE_COLUMNS = [
    Column("File", Kind.TEXT, lambda r: r.name),
    Column("Status", Kind.TEXT, _status),
    Column("Reasons", Kind.TEXT, _reasons),
    Column("Reason codes", Kind.TEXT, _codes),
    Column("Type", Kind.TEXT, lambda r: r.invoice.doc_type.value),
    Column("Supplier", Kind.TEXT, lambda r: r.invoice.supplier_name),
    Column("Supplier tax ID", Kind.TEXT, lambda r: r.invoice.supplier_tax_id),
    Column("Invoice number", Kind.TEXT, lambda r: r.invoice.invoice_number),
    Column("Invoice date", Kind.DATE, lambda r: r.invoice.invoice_date),
    Column("Due date", Kind.DATE, lambda r: r.invoice.due_date),
    Column("Currency", Kind.TEXT, lambda r: r.invoice.currency),
    Column("PO number", Kind.TEXT, lambda r: r.invoice.po_number),
    Column("Subtotal", Kind.MONEY, lambda r: r.invoice.subtotal),
    Column("Tax", Kind.MONEY, lambda r: r.invoice.tax),
    Column("Discount", Kind.MONEY, lambda r: r.invoice.discount),
    Column("Shipping", Kind.MONEY, lambda r: r.invoice.shipping),
    Column("Total", Kind.MONEY, lambda r: r.invoice.total),
    Column("Lines", Kind.NUMBER, lambda r: len(r.invoice.line_items)),
    Column("Model", Kind.TEXT, lambda r: r.invoice.model),
    Column("Processed", Kind.DATETIME, lambda r: r.processed_at),
]

LINE_COLUMNS = [
    Column("File", Kind.TEXT, lambda r, i, line: r.name),
    Column("Supplier", Kind.TEXT, lambda r, i, line: r.invoice.supplier_name),
    Column("Invoice number", Kind.TEXT, lambda r, i, line: r.invoice.invoice_number),
    Column("Line", Kind.NUMBER, lambda r, i, line: i),
    Column("Description", Kind.TEXT, lambda r, i, line: line.description),
    Column("SKU", Kind.TEXT, lambda r, i, line: line.sku),
    Column("Quantity", Kind.NUMBER, lambda r, i, line: line.quantity),
    Column("Unit price", Kind.MONEY, lambda r, i, line: line.unit_price),
    Column("Line total", Kind.MONEY, lambda r, i, line: line.line_total),
    Column("Tax rate %", Kind.NUMBER, lambda r, i, line: line.tax_rate),
    Column("Currency", Kind.TEXT, lambda r, i, line: r.invoice.currency),
    Column("Status", Kind.TEXT, lambda r, i, line: _status(r)),
]

REVIEW_COLUMNS = [
    Column("File", Kind.TEXT, lambda r: r.name),
    Column("Status", Kind.TEXT, _status),
    Column("Reasons", Kind.TEXT, _reasons),
    Column("Supplier", Kind.TEXT, lambda r: r.invoice.supplier_name),
    Column("Invoice number", Kind.TEXT, lambda r: r.invoice.invoice_number),
    Column("Invoice date", Kind.DATE, lambda r: r.invoice.invoice_date),
    Column("Currency", Kind.TEXT, lambda r: r.invoice.currency),
    Column("Total", Kind.MONEY, lambda r: r.invoice.total),
    Column("PDF", Kind.TEXT, lambda r: r.pdf_path),
    Column("Processed", Kind.DATETIME, lambda r: r.processed_at),
]


def _needs_review(invoice: Invoice) -> bool:
    return invoice.status != Status.OK


def build_sheets(results: Iterable[Result]) -> list[Sheet]:
    """The three sheets: Invoices, Line items, Needs review."""
    results = list(results)
    invoices = Sheet("Invoices", INVOICE_COLUMNS, [], set())
    lines = Sheet("Line items", LINE_COLUMNS, [], set())
    review = Sheet("Needs review", REVIEW_COLUMNS, [], set())
    for r in results:
        flagged = _needs_review(r.invoice)
        if flagged:
            invoices.review_rows.add(len(invoices.rows))
        invoices.rows.append([c.value(r) for c in INVOICE_COLUMNS])
        for i, line in enumerate(r.invoice.line_items, start=1):
            if flagged:
                lines.review_rows.add(len(lines.rows))
            lines.rows.append([c.value(r, i, line) for c in LINE_COLUMNS])
        if flagged:
            review.review_rows.add(len(review.rows))
            review.rows.append([c.value(r) for c in REVIEW_COLUMNS])
    return [invoices, lines, review]


# --- CSV ------------------------------------------------------------------------------------


def _text(value: object) -> str:
    """A cell as exact text: Decimal without exponent, ISO dates, "" for empty."""
    if value is None:
        return ""
    if isinstance(value, Decimal):
        return format(value, "f")
    if isinstance(value, datetime):
        return value.isoformat(timespec="seconds")
    if isinstance(value, date):
        return value.isoformat()
    return str(value)


def _safe(text: str) -> str:
    """Text that can't run as a formula in Excel or Sheets."""
    return "'" + text if text.startswith(FORMULA_START) else text


def write_csv(path: Path, sheet: Sheet) -> None:
    def rows():
        yield sheet.headers
        for row in sheet.rows:
            yield [
                _safe(_text(v)) if c.kind == Kind.TEXT else _text(v)
                for c, v in zip(sheet.columns, row, strict=True)
            ]

    # utf-8-sig so Excel shows £ and € correctly when it opens the file.
    _atomic_write(path, lambda f: csv.writer(f).writerows(rows()), encoding="utf-8-sig")


# --- Excel ----------------------------------------------------------------------------------

_FORMATS = {
    Kind.MONEY: "#,##0.00##",
    Kind.NUMBER: "#,##0.####",
    Kind.DATE: "yyyy-mm-dd",
    Kind.DATETIME: "yyyy-mm-dd hh:mm",
}


def write_excel(path: Path, sheets: list[Sheet]) -> None:
    book = Workbook()
    book.remove(book.active)
    for sheet in sheets:
        ws = book.create_sheet(sheet.name)
        ws.append(sheet.headers)
        for cell in ws[1]:
            cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        for index, row in enumerate(sheet.rows):
            excel_row = index + 2
            for col, (column, value) in enumerate(zip(sheet.columns, row, strict=True), start=1):
                cell = ws.cell(row=excel_row, column=col)
                if index in sheet.review_rows:
                    cell.fill = REVIEW_FILL
                if value is None:
                    continue
                if column.kind == Kind.TEXT:
                    cell.value = str(value)
                    cell.data_type = "s"  # never a formula, even if it starts with "="
                else:
                    cell.value = value
                    cell.number_format = _FORMATS[column.kind]
        ws.freeze_panes = "A2"
        if sheet.rows:
            ws.auto_filter.ref = ws.dimensions
        for col, column in enumerate(sheet.columns, start=1):
            longest = max([len(column.header)] + [len(_text(r[col - 1])) for r in sheet.rows])
            ws.column_dimensions[get_column_letter(col)].width = min(max(longest + 2, 10), 60)
    _atomic_write(path, book.save, binary=True)


# --- Google Sheets --------------------------------------------------------------------------


def sheets_payload(sheets: list[Sheet]) -> dict[str, list[list]]:
    """{sheet name: rows} for the Sheets API with USER_ENTERED input: numbers as numbers, dates
    as ISO text (which Sheets reads as dates in any locale), other text kept from becoming a
    formula."""

    def cell(column: Column, value: object) -> object:
        if value is None:
            return ""
        if column.kind in (Kind.MONEY, Kind.NUMBER):
            return float(value) if isinstance(value, Decimal) else value
        if column.kind in (Kind.DATE, Kind.DATETIME):
            return _text(value).replace("T", " ")
        return _safe(str(value))

    return {
        s.name: [s.headers]
        + [[cell(c, v) for c, v in zip(s.columns, row, strict=True)] for row in s.rows]
        for s in sheets
    }


def spreadsheet_key(id_or_url: str) -> str:
    """The key from a spreadsheet ID or its full URL."""
    match = re.search(r"/spreadsheets/d/([A-Za-z0-9_-]+)", id_or_url)
    return match.group(1) if match else id_or_url.strip()


def write_google_sheet(spreadsheet: str, payload: dict[str, list[list]], credentials: Path) -> str:
    """Replace each named worksheet's contents with `payload`. Returns the spreadsheet's URL."""
    import gspread  # optional dependency: uv sync --extra sheets

    book = gspread.service_account(filename=str(credentials)).open_by_key(
        spreadsheet_key(spreadsheet)
    )
    existing = {ws.title: ws for ws in book.worksheets()}
    for name, rows in payload.items():
        ws = existing.get(name) or book.add_worksheet(
            title=name, rows=max(len(rows), 1), cols=len(rows[0])
        )
        ws.clear()
        ws.update(rows, "A1", value_input_option="USER_ENTERED")
        ws.freeze(rows=1)
        ws.format("1:1", {"textFormat": {"bold": True}})
    return book.url


# --- all outputs ----------------------------------------------------------------------------


def _atomic_write(path: Path, write: Callable, *, binary: bool = False, encoding: str = "utf-8"):
    """Write to a temporary file next to `path`, then swap it in, so a crash or a file open in
    Excel never leaves a half-written export."""
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, tmp = tempfile.mkstemp(dir=path.parent, prefix=f".{path.name}.", suffix=".tmp")
    try:
        if binary:
            os.close(fd)
            write(tmp)
        else:
            with os.fdopen(fd, "w", encoding=encoding, newline="") as f:
                write(f)
        os.replace(tmp, path)
    except BaseException:
        Path(tmp).unlink(missing_ok=True)
        raise
